"""Markdown to PDF."""
from __future__ import annotations

import os
import re
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CJK_FONT = "Noto Sans CJK SC"
LATIN_FONT = "DejaVu Sans"
MONO_FONT = "DejaVu Sans Mono"


def _set_cjk(run_or_style, east=CJK_FONT, latin=LATIN_FONT):
    """Set both latin and east-asian font names on a style or run."""
    el = run_or_style._element
    rPr = el.get_or_add_rPr() if hasattr(el, "get_or_add_rPr") else None
    if rPr is None:
        return
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)
    rFonts.set(qn("w:eastAsia"), east)
    rFonts.set(qn("w:cs"), latin)


def _style_fonts(doc):
    for name, size, bold, color in (
        ("Normal", 10.5, False, None),
        ("Title", 17, True, RGBColor(0x1F, 0x3A, 0x5F)),
        ("Heading 1", 15, True, RGBColor(0x1F, 0x3A, 0x5F)),
        ("Heading 2", 12.5, True, RGBColor(0x1F, 0x3A, 0x5F)),
        ("Heading 3", 11, True, RGBColor(0x33, 0x33, 0x33)),
        ("Caption", 9, False, RGBColor(0x44, 0x44, 0x44)),
    ):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = LATIN_FONT
        st.font.size = Pt(size)
        st.font.bold = bold
        if color is not None:
            st.font.color.rgb = color
        _set_cjk(st)
        pf = st.paragraph_format
        pf.space_before = Pt(6 if name.startswith("Heading") else 0)
        pf.space_after = Pt(4)
        if name == "Normal":
            pf.line_spacing = 1.28
    # Title centred
    try:
        doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    except KeyError:
        pass


def _add_page_number_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    run.font.size = Pt(9)
    run.font.name = LATIN_FONT


def _add_runs(par, text):
    """Parse **bold**, *italic* and `code` into runs."""
    tokens = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
    for tok in tokens:
        if not tok:
            continue
        bold = italic = mono = False
        if tok.startswith("**") and tok.endswith("**"):
            tok, bold = tok[2:-2], True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            tok, italic = tok[1:-1], True
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            tok, mono = tok[1:-1], True
        run = par.add_run(tok)
        run.bold = bold
        run.italic = italic
        if mono:
            run.font.name = MONO_FONT
            run.font.size = Pt(9.5)


def _shade_cell(cell, hexcolor="EAF0F7"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _add_table(doc, rows):
    header, body = rows[0], rows[1:]
    ncol = len(header)
    table = doc.add_table(rows=1 + len(body), cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for j, cell_text in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_runs(p, cell_text)
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(8.5)
        _shade_cell(cell)
    for i, row in enumerate(body, start=1):
        for j in range(ncol):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            txt = row[j] if j < len(row) else ""
            _add_runs(p, txt)
            for r in p.runs:
                r.font.size = Pt(8.5)
    # spacing after table
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def _add_image(doc, path, caption, width_cm=16.2):
    if not os.path.exists(path):
        p = doc.add_paragraph()
        p.add_run(f"[缺失图片: {path}]").italic = True
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Cm(width_cm))
    if caption:
        cp = doc.add_paragraph(style="Caption")
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_runs(cp, caption)


def _parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        line = lines[i].strip()
        if re.match(r"^\|?\s*:?-{2,}", line):
            i += 1
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        rows.append(cells)
        i += 1
    return rows, i


def markdown_to_docx(md_path, docx_path, title=None):
    base_dir = os.path.dirname(os.path.abspath(md_path))
    with open(md_path, "r", encoding="utf-8") as f:
        text = f.read()
    lines = text.splitlines()
    doc = Document()
    _style_fonts(doc)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    _add_page_number_footer(sec)
    if title:
        doc.add_paragraph(title, style="Title")
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s == "<!-- pagebreak -->":
            doc.add_page_break()
            i += 1
            continue
        if s.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            r = p.add_run("\n".join(code))
            r.font.name = MONO_FONT
            r.font.size = Pt(8.5)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            level = len(m.group(1))
            h = doc.add_heading(level=min(level, 4))
            _add_runs(h, m.group(2))
            if level == 1:
                h.alignment = WD_ALIGN_PARAGRAPH.LEFT
            i += 1
            continue
        if s.startswith("!["):
            m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", s)
            if m:
                img_path = m.group(2)
                if not os.path.isabs(img_path):
                    img_path = os.path.join(base_dir, img_path)
                _add_image(doc, img_path, m.group(1))
                i += 1
                continue
        if s.startswith("EQUATION:"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(s[len("EQUATION:"):].strip())
            r.italic = True
            r.font.size = Pt(10)
            i += 1
            continue
        if s.startswith("NOTE:"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            r = p.add_run(s[len("NOTE:"):].strip())
            r.italic = True
            r.font.size = Pt(9)
            i += 1
            continue
        if s.startswith("|"):
            rows, i = _parse_table(lines, i)
            if rows:
                _add_table(doc, rows)
            continue
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            r = p.add_run(s.lstrip("> ").strip())
            r.italic = True
            r.font.size = Pt(9.5)
            i += 1
            continue
        if re.match(r"^[-*]\s+", s):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, re.sub(r"^[-*]\s+", "", s))
            i += 1
            continue
        if re.match(r"^\d+[.)]\s+", s):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.45)
            _add_runs(p, s)          # keep the original number, no auto numbering
            i += 1
            continue
        if s == "---":
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "999999")
            pbdr.append(bottom)
            pPr.append(pbdr)
            i += 1
            continue
        # paragraph: gather continuation lines until blank/block start
        para = [s]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("#", "|", "![", ">", "---", "```"))
                    or re.match(r"^[-*]\s+", nxt) or re.match(r"^\d+[.)]\s+", nxt)):
                break
            para.append(nxt)
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _add_runs(p, _join_para(para))
    doc.save(docx_path)
    return docx_path


def _join_para(parts):
    """Join wrapped source lines: no space after CJK characters/punctuation."""
    out = parts[0]
    for nxt in parts[1:]:
        if re.search(r"[\u4e00-\u9fff，。；：、）】》]$", out):
            out += nxt
        else:
            out += " " + nxt
    return out


def docx_to_pdf(docx_path, pdf_path, outdir=None):
    outdir = outdir or os.path.dirname(os.path.abspath(pdf_path)) or "."
    os.makedirs(outdir, exist_ok=True)
    cmd = ["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir,
           os.path.abspath(docx_path)]
    res = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
    produced = os.path.join(outdir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
    if not os.path.exists(produced):
        raise RuntimeError(f"LibreOffice PDF conversion failed: {res.stderr}\n{res.stdout}")
    if os.path.abspath(produced) != os.path.abspath(pdf_path):
        os.replace(produced, pdf_path)
    return pdf_path


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return 1
    md_path, pdf_path = sys.argv[1], sys.argv[2]
    title = sys.argv[3] if len(sys.argv) > 3 else None
    tmp_docx = os.path.splitext(pdf_path)[0] + ".docx"
    markdown_to_docx(md_path, tmp_docx, title=title)
    docx_to_pdf(tmp_docx, pdf_path)
    print("PDF:", pdf_path)
    print("DOCX (intermediate):", tmp_docx)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
